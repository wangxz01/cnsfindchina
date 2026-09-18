"""Excel output: summary plus one sheet per issue, with native date cells."""
from __future__ import annotations

import os
import re
import tempfile
import time
from datetime import date
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from article_metadata import china_label, FAILED_TITLES

BASE_COLUMNS = [
    ("相关网址", "url"), ("标题", "title"), ("DOI", "doi"), ("类型", "type"),
    ("第一作者", "first_author"), ("第一作者单位（未核实见状态列）", "first_aff"),
    ("国家", "first_author_country"), ("是否中国", "is_china_label"), ("所有作者", "authors_joined"),
]
QUALITY_COLUMNS = [("提取状态", "extraction_status"), ("待补字段", "missing_fields"), ("日期来源", "date_evidence")]
CELL_COLUMNS = BASE_COLUMNS + [("Available online", "available_online"), ("Version of Record", "version_of_record")] + QUALITY_COLUMNS
NATURE_COLUMNS = BASE_COLUMNS + [("发布日期", "published_date")] + QUALITY_COLUMNS
DEFAULT_COLUMNS = NATURE_COLUMNS
DATE_FIELDS = {"available_online", "version_of_record", "published_date"}


def column_widths(columns):
    widths = {"url": 55, "title": 55, "doi": 25, "type": 18, "first_author": 22,
              "first_aff": 65, "first_author_country": 22, "is_china_label": 12,
              "authors_joined": 50, "extraction_status": 15, "missing_fields": 40, "date_evidence": 55}
    return [widths.get(key, 19) for _, key in columns]


def _sheet_name_from_url(url):
    m = re.search(r"/vol(?:umes)?/(\d+)/issues?/(\d+)|/toc/science/(\d+)/(\d+)", url)
    if m:
        volume, issue = (m.group(1), m.group(2)) if m.group(1) else (m.group(3), m.group(4))
        return f"v{volume}-i{issue}"
    return re.sub(r"[\\/*?:\[\]]", "-", "-".join(url.rstrip("/").split("/")[-2:]))[:31] or "issue"


def _cell_value(key, fields):
    if callable(key):
        return key(fields)
    if key == "authors_joined":
        return "; ".join(fields.get("authors", []))
    if key == "is_china_label":
        return china_label(fields.get("is_china"))
    value = fields.get(key, "")
    if key in DATE_FIELDS:
        try:
            return date.fromisoformat(value) if value else None
        except (ValueError, TypeError):
            return None
    if key == "extraction_status":
        return {"complete": "完整", "partial": "待补全", "failed": "抓取失败"}.get(value, "待核实")
    if isinstance(value, dict):
        return "; ".join(f"{k}: {v}" for k, v in value.items())
    if isinstance(value, list):
        return "; ".join(map(str, value))
    return value


def _append(ws, values):
    ws.append(values)
    for cell in ws[ws.max_row]:
        # External titles/affiliations must be written as text, never Excel formulas.
        if isinstance(cell.value, str):
            cell.data_type = 's'
        if isinstance(cell.value, date):
            cell.number_format = 'yyyy-mm-dd'
        cell.alignment = Alignment(vertical='top', wrap_text=True)


def write_one_sheet(wb, issue_url, results, columns=DEFAULT_COLUMNS, col_widths=None):
    ws = wb.create_sheet(_sheet_name_from_url(issue_url))
    _append(ws, [issue_url])
    _append(ws, [name for name, _ in columns])
    for cell in ws[2]:
        cell.font = Font(bold=True)
    last_section = None
    for section, url, fields in results:
        if section != last_section:
            _append(ws, [section])
            last_section = section
        fields = dict(fields, url=url)
        _append(ws, [_cell_value(key, fields) for _, key in columns])
        if fields.get('is_china') is True:
            for cell in ws[ws.max_row]:
                cell.fill = PatternFill('solid', fgColor='FFF0F0')
    for i, width in enumerate(col_widths or column_widths(columns), 1):
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = 'A3'


def write_summary(wb, all_issues, issue_meta):
    ws = wb.active
    ws.title = '汇总'
    _append(ws, ['统计口径：署名第一作者；任一明确关联单位在中国则计入；未知不计为否。'])
    headers = ['Issue URL', '计划篇数', '已取得记录', '成功访问', '抓取失败', '中国', '非中国', '国家待核实', '数据待补全', '数量校验', '异常说明']
    _append(ws, headers)
    for url, results in all_issues:
        fields = [f for _, _, f in results]
        success = [f for f in fields if f.get('title') and f['title'] not in FAILED_TITLES]
        meta = issue_meta.get(url, {})
        check = meta.get('count_check')
        check_label = ('通过' if check.get('matched') else f"不一致 {check['actual']}/{check['expected']}") if check else '未校验'
        _append(ws, [url, meta.get('expected'), len(fields), len(success), len(fields)-len(success),
                     sum(f.get('is_china') is True for f in success), sum(f.get('is_china') is False for f in success),
                     sum(f.get('is_china') is None for f in fields),
                     sum(f.get('extraction_status') != 'complete' for f in fields), check_label, meta.get('error', '')])
    ws.column_dimensions['A'].width = 65
    for i in range(2, len(headers)+1):
        ws.column_dimensions[get_column_letter(i)].width = 20
    ws.freeze_panes = 'B3'
    ws.auto_filter.ref = f'A2:K{max(2, ws.max_row)}'


def write_excel(out_path, all_issues, columns=DEFAULT_COLUMNS, col_widths=None, issue_meta=None):
    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.Workbook()
    write_summary(wb, all_issues, issue_meta or {})
    for issue_url, results in all_issues:
        write_one_sheet(wb, issue_url, results, columns, col_widths)
    fd, temp = tempfile.mkstemp(dir=out_path.parent, suffix='.xlsx')
    os.close(fd)
    try:
        wb.save(temp)
        try:
            os.replace(temp, out_path)
        except PermissionError:
            out_path = out_path.with_name(f'{out_path.stem}.{time.time_ns()}.xlsx')
            os.replace(temp, out_path)
            print(f'[warn] 文件被占用，改写到: {out_path}')
        return str(out_path)
    finally:
        wb.close()
        if os.path.exists(temp):
            os.unlink(temp)
