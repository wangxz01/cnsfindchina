"""Excel 输出（方案 A 布局）。

布局：
  每个 issue 一个 sheet（按 vol/issue 命名）
  Row 1, Col A                 = issue URL
  之后按 section 出现顺序：
    分类名独占一行（Col A）
    每篇文章一行：[URL | 标题 | DOI | 作者(分号拼接) | 首条 affiliation]
"""
import re
import time
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter


def _sheet_name_from_url(url: str) -> str:
    m = re.search(r"/vol/(\d+)/issue/(\d+)", url)
    if m:
        return f"v{m.group(1)}-i{m.group(2)}"
    parts = url.rstrip("/").split("/")
    name = "-".join(parts[-2:]) or "issue"
    return name[:31]  # Excel sheet 名最长 31 字符


def write_one_sheet(wb, issue_url: str, results: list) -> None:
    ws = wb.create_sheet(_sheet_name_from_url(issue_url))
    ws.append([issue_url])
    last_section = None
    for section, url, f in results:
        if section != last_section:
            ws.append([section])
            last_section = section
        ws.append([
            url,
            f.get("title", ""),
            f.get("doi", ""),
            "; ".join(f.get("authors", [])),
            f.get("first_aff", ""),
        ])
    widths = [60, 60, 30, 40, 60]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def write_excel(out_path: str, all_issues: list) -> None:
    """all_issues: [(issue_url, results), ...]。每个 issue 一个 sheet。"""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)  # 删除默认 sheet
    for issue_url, results in all_issues:
        write_one_sheet(wb, issue_url, results)

    try:
        wb.save(out_path)
    except PermissionError:
        # 文件被占用（如 Excel 打开中）→ 写到带时间戳的备用文件
        ts = time.strftime("%Y%m%d_%H%M%S")
        alt = str(Path(out_path).with_suffix(f".{ts}.xlsx"))
        print(f"[warn] {out_path} 被占用，改写到: {alt}")
        wb.save(alt)

