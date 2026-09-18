"""Offline regression tests. Fixtures are synthetic and contain no real paper data."""
from __future__ import annotations

import json
import sys
import tempfile
import threading
import unittest
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'src'))

import openpyxl
import scraper
import nature_scraper
import science_scraper
import scraper_common as common
import unified_web as web
from article_metadata import (publication_dates, normalize_date, enrich_fields,
                              expected_article_ids, count_check, finalize_fields)
from countries import parse_country
from excel_writer import write_excel, CELL_COLUMNS, NATURE_COLUMNS, _cell_value

FIXTURES = Path(__file__).parent / 'fixtures'


def fixture(name):
    return (FIXTURES / name).read_text(encoding='utf-8')


def complete_fields(source='cell'):
    fields = dict(title='Fixture paper', doi='10.1016/j.cell.fixture', first_author='Alice Li',
                  authors=['Alice Li'], first_aff='Institute, Beijing, China', is_china=True,
                  first_author_country='China', affiliation_verified=True,
                  available_online='2026-05-02', version_of_record='2026-05-10',
                  published_date='2026-05-02', date_evidence={}, url='https://example.org/article')
    return finalize_fields(fields, source)


class MetadataTests(unittest.TestCase):
    def test_cell_distinct_dates_and_not_received_or_issue_date(self):
        dates = publication_dates(fixture('cell_history.html'), 'cell')
        self.assertEqual(dates['available_online'], '2026-05-02')
        self.assertEqual(dates['version_of_record'], '2026-05-10')
        self.assertEqual(dates['published_date'], '')
        self.assertIn('Available online', dates['date_evidence']['available_online'])

    def test_cell_missing_date_never_copied_from_other_date(self):
        dates = publication_dates('<meta name="citation_publication_date" content="2026/05/20"><p>Available online 2 May 2026</p>', 'cell')
        self.assertEqual(dates['version_of_record'], '')

    def test_dates_support_metadata_attribute_order_and_iso(self):
        html = "<meta content='2026/5/2' name='citation_online_date'>"
        for source in ('nature', 'science'):
            self.assertEqual(publication_dates(html, source)['published_date'], '2026-05-02')
        self.assertEqual(normalize_date('2026-05-02T08:00:00Z'), '2026-05-02')
        self.assertEqual(normalize_date('Feb 31, 2026'), '')

    def test_explicit_vor_key_and_label(self):
        html = '<script>{"vor-available-online-date":"2026-05-10"}</script>'
        self.assertEqual(publication_dates(html, 'cell')['version_of_record'], '2026-05-10')
        dates = publication_dates('Version of Record available online 10 May 2026', 'cell')
        self.assertEqual(dates['version_of_record'], '2026-05-10')
        self.assertEqual(dates['available_online'], '')

    def test_cell_current_dates_object(self):
        html = '<script>window.data={"availableOnlineDate":"2026-05-02","dates":{"Available online":"2 May 2026","Revised":[],"Publication date":"20 May 2026","Version of Record":"10 May 2026"}}</script>'
        dates = publication_dates(html, 'cell')
        self.assertEqual(dates['available_online'], '2026-05-02')
        self.assertEqual(dates['version_of_record'], '2026-05-10')
        self.assertEqual(dates['date_evidence']['version_of_record'], 'dates.Version of Record: 10 May 2026')

    def test_dates_ignore_reference_dates(self):
        html = '<section id="references">Available online 1 January 1999</section><p>Available online 2 May 2026; Version of Record 10 May 2026</p>'
        self.assertEqual(publication_dates(html, 'cell')['available_online'], '2026-05-02')

    def test_country_aliases_normalize_accents_and_punctuation(self):
        self.assertEqual(parse_country('Institute, Ankara, Türkiye'), 'Turkey')
        self.assertEqual(parse_country('Institute, Beijing, P.R. China'), 'China')

    def test_cell_reads_dates_after_show_more_and_history(self):
        class Locator:
            def count(self): return 1
            def click(self, **kwargs): page.stage = 2
        class Page:
            stage = 0
            def content(self):
                return fixture('cell_article.html') + (fixture('cell_history.html') if self.stage == 2 else '')
            def get_by_role(self, *args, **kwargs): return Locator()
            def wait_for_timeout(self, *args): pass
            def query_selector(self, *args): return None
        page = Page()
        def show_more(p): p.stage = 1; return True
        with patch.object(scraper, '_click_show_more_js', side_effect=show_more) as click, patch.object(scraper, 'human_pause'), patch.object(scraper, 'random_mouse_jitter'):
            fields = scraper.extract_fields(page, 'https://example.org/article')
        click.assert_called_once()
        self.assertEqual(fields['available_online'], '2026-05-02')
        self.assertEqual(fields['version_of_record'], '2026-05-10')
        self.assertIs(fields['is_china'], True)
        self.assertEqual(fields['first_author_country'], 'China')
        self.assertEqual(fields['extraction_status'], 'complete')

    def test_cell_uses_linked_aff2_not_first_aff1(self):
        fields = scraper._extract_from_html(fixture('cell_article.html'))
        self.assertIn('USA', fields['first_aff'])
        fields = enrich_fields(fields, fixture('cell_article.html'), 'cell', parse_country)
        self.assertEqual(fields['first_aff'], 'Institute, Beijing, China')
        self.assertIs(fields['is_china'], True)

    def test_cell_decodes_author_names_before_matching(self):
        html = fixture('cell_article.html').replace('Alice', r'Alic\u00e9')
        fields = scraper._extract_from_html(html)
        fields = enrich_fields(fields, html, 'cell', parse_country)
        self.assertEqual(fields['first_author'], 'Alicé Li')
        self.assertIs(fields['is_china'], True)

    def test_nature_first_author_multiple_affiliations(self):
        html = fixture('nature_article.html')
        fields = nature_scraper._extract_fields_from_html(html, 'https://example.org/article')
        fields = enrich_fields(fields, html, 'nature', parse_country)
        self.assertEqual(fields['first_author'], 'Alice Li')
        self.assertEqual(fields['first_author_affiliations'], ['Institute, Boston, USA', 'Institute, Beijing, China'])
        self.assertIs(fields['is_china'], True)
        self.assertEqual(fields['published_date'], '2026-05-02')

    def test_unlinked_affiliation_is_unknown(self):
        html = fixture('nature_article.html').replace('Alice Li', 'Other Person')
        fields = enrich_fields(dict(first_author='Alice Li', authors=['Alice Li']), html, 'nature', parse_country)
        self.assertIsNone(fields['is_china'])
        self.assertIn('author_affiliation_mapping', fields['missing_fields'])
        self.assertEqual(_cell_value('is_china_label', fields), '待核实')

    def test_science_contributor_boundary(self):
        html = fixture('science_article.html')
        fields = science_scraper._extract_fields_from_html(html, 'https://www.science.org/doi/10.1126/science.fixture')
        fields = enrich_fields(fields, html, 'science', parse_country)
        self.assertIs(fields['is_china'], False)
        self.assertEqual(fields['first_author_country'], 'USA')
        self.assertEqual(fields['published_date'], '2026-05-02')
        missing = html.replace('<div property="affiliation"><span property="name">Institute, Boston, USA</span></div>', '')
        fields = enrich_fields(dict(first_author='Alice Li'), missing, 'science', parse_country)
        self.assertIsNone(fields['is_china'])

    def test_unknown_country_is_not_no(self):
        html = fixture('science_article.html').replace('Boston, USA', 'Unknown Place')
        fields = enrich_fields(dict(first_author='Alice Li'), html, 'science', parse_country)
        self.assertIsNone(fields['is_china'])

    def test_count_excludes_corrections_and_related(self):
        ids = expected_article_ids(fixture('nature_issue.html'), 'nature')
        self.assertEqual(ids, {'s41586-026-12345-6'})
        targets = [('Articles', 'https://www.nature.com/articles/s41586-026-12345-6', 'Paper')]
        self.assertTrue(count_check(fixture('nature_issue.html'), targets, 'nature', 'issue')['matched'])

    def test_same_count_different_ids_is_mismatch(self):
        targets = [('Articles', 'https://www.nature.com/articles/s41586-026-99999-9', 'Other')]
        check = count_check(fixture('nature_issue.html'), targets, 'nature', 'issue')
        self.assertEqual(check['actual'], check['expected'])
        self.assertFalse(check['matched'])
        self.assertEqual(check['missing_ids'], ['s41586-026-12345-6'])

    def test_zero_zero_is_not_success(self):
        self.assertFalse(count_check('<main></main>', [], 'cell', 'issue')['matched'])

    def test_cf_final_retry_can_succeed_and_pause_receives_page(self):
        good = complete_fields()
        page = object()
        with patch.object(scraper, 'is_cloudflare', return_value=False), patch.object(scraper, 'wait_until_cf_clear', return_value=True), patch.object(scraper, 'human_pause') as pause:
            values = iter([{'title':'Just a moment'}, good])
            result = scraper.extract_with_cf_retry(page, 'url', scraper.ScraperCallbacks(), 'Article',
                      lambda *args: next(values), pause, max_retries=1)
        self.assertIs(result, good)
        pause.assert_called_once_with(page, 1.2, 2.8)


class StorageTests(unittest.TestCase):
    def test_partial_old_corrupt_cache_miss_and_valid_hit(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory)
            fields = complete_fields()
            common.save_to_cache(path, 'article', fields)
            self.assertIsNotNone(common.load_from_cache(path, 'article'))
            fields['version_of_record'] = ''
            common.save_to_cache(path, 'article', fields)
            self.assertIsNone(common.load_from_cache(path, 'article'))
            common.cache_path(path, 'article').write_text('{bad')
            self.assertIsNone(common.load_from_cache(path, 'article'))
            common.cache_path(path, 'article').write_text(json.dumps(complete_fields()))
            self.assertIsNone(common.load_from_cache(path, 'article'))

    def test_retry_preserves_verified_date(self):
        with tempfile.TemporaryDirectory() as directory:
            old = complete_fields(); old['version_of_record'] = ''
            common.save_to_cache(Path(directory), 'article', old)
            new = complete_fields(); new['available_online'] = ''
            merged = common.merge_cached_fields(Path(directory), 'article', new)
            self.assertEqual(merged['available_online'], '2026-05-02')
            self.assertEqual(merged['version_of_record'], '2026-05-10')
            self.assertEqual(merged['extraction_status'], 'complete')

    def test_excel_dates_summary_unknown_and_formula_as_text(self):
        with tempfile.TemporaryDirectory() as directory:
            fields = complete_fields(); fields.update(title='=1+1', is_china=None)
            out = write_excel(str(Path(directory)/'nested'/'result.xlsx'), [('https://example.org/vol/1/issue/2', [('Articles', fields['url'], fields)])], columns=CELL_COLUMNS)
            wb = openpyxl.load_workbook(out)
            try:
                ws = wb['v1-i2']
                header = {cell.value:cell.column for cell in ws[2]}
                online = ws.cell(4, header['Available online'])
                vor = ws.cell(4, header['Version of Record'])
                self.assertEqual(online.value, datetime(2026,5,2))
                self.assertEqual(vor.value, datetime(2026,5,10))
                self.assertEqual(online.number_format, 'yyyy-mm-dd')
                self.assertEqual(ws.cell(4,header['是否中国']).value, '待核实')
                self.assertEqual(ws.cell(4,header['标题']).data_type, 's')
                self.assertEqual(wb['汇总'].cell(3,8).value, 1)
            finally: wb.close()

    def test_other_sources_have_one_date_column(self):
        self.assertEqual([key for _,key in NATURE_COLUMNS if 'date' in key or key in ('available_online','version_of_record')], ['published_date','date_evidence'])

    def test_empty_export_and_locked_destination(self):
        with tempfile.TemporaryDirectory() as directory:
            target = str(Path(directory)/'result.xlsx')
            actual_replace = common.os.replace
            calls = []
            def replace(src, dst):
                calls.append(str(dst))
                if len(calls) == 1: raise PermissionError()
                return actual_replace(src, dst)
            with patch('excel_writer.os.replace', side_effect=replace):
                out = write_excel(target, [])
            self.assertNotEqual(out, target)
            self.assertTrue(Path(out).exists())


class LifecycleTests(unittest.TestCase):
    def run_case(self, module, cancelled=False, partial=False):
        class Callback(scraper.ScraperCallbacks):
            cancelled = False
            events = []
            def log(self, msg): pass
            def is_cancelled(self): return self.cancelled
            def on_state(self, state): self.events.append(state)
        cb = Callback()
        context = SimpleNamespace(new_page=lambda:object(), close=lambda:None)
        class PW:
            def __enter__(self): return SimpleNamespace(chromium=SimpleNamespace(launch_persistent_context=lambda **k:context))
            def __exit__(self,*args): pass
        def process(page,url,**kwargs):
            fields=complete_fields('cell' if module is scraper else 'nature')
            if partial: fields['published_date']=fields['available_online']=''
            callback=kwargs['cb']
            callback.on_state(dict(phase='issue_plan',targets=[('Articles','article','title')]))
            callback.on_state(dict(phase='count_check',issue_url=url,actual=1,expected=1,matched=True))
            callback.on_state(dict(phase='article_done',url='article',fields=fields))
            cb.cancelled=cancelled
            return [('Articles','article',fields)]
        with tempfile.TemporaryDirectory() as directory, patch.object(module,'sync_playwright',PW), patch.object(module,'process_issue',side_effect=process):
            module.run_scraper(['https://example.org/vol/1/issue/1'],Path(directory)/'out.xlsx',cb=cb)
        return cb.events[-1]['phase']

    def test_cancellation_stays_cancelled_all_sources(self):
        for module in (scraper,nature_scraper,science_scraper):
            with self.subTest(module=module.__name__):
                self.assertEqual(self.run_case(module,cancelled=True),'cancelled')

    def test_partial_and_complete_are_distinct(self):
        self.assertEqual(self.run_case(nature_scraper,partial=True),'partial_done')
        self.assertEqual(self.run_case(nature_scraper),'all_done')

    def test_stop_interrupts_delay(self):
        callback = SimpleNamespace(is_cancelled=lambda:True)
        with self.assertRaises(common.ScrapeCancelled):
            common.cancellable_sleep(60, callback)


class WebTests(unittest.TestCase):
    def setUp(self):
        self.sessions_patch=patch.object(web,'SESSIONS',{k:web.Session(k) for k in web.SOURCES})
        self.bus_patch=patch.object(web,'bus',web.EventBus())
        self.sessions_patch.start();self.bus_patch.start()
        self.addCleanup(self.sessions_patch.stop);self.addCleanup(self.bus_patch.stop)

    def test_snapshot_recovers_articles_logs_dates_and_check(self):
        cb=web.WebCallbacks('cell')
        cb.log('test log')
        cb.on_state(dict(phase='article_done',url='article',fields=complete_fields()))
        cb.on_state(dict(phase='count_check',issue_url='issue',actual=1,expected=1,matched=True))
        q=web.bus.subscribe()
        event=q.get_nowait()
        self.assertEqual(event['type'],'snapshot')
        self.assertEqual(event['data']['articles'][0]['version_of_record'],'2026-05-10')
        self.assertEqual(event['data']['logs'],['test log'])
        self.assertTrue(event['data']['count_checks']['issue']['matched'])
        self.assertEqual(web.status('cell')['articles'],event['data']['articles'])
        web.bus.unsubscribe(q)

    def test_duplicate_event_does_not_duplicate_result(self):
        cb=web.WebCallbacks('cell')
        event=dict(phase='article_done',url='article',fields=complete_fields())
        cb.on_state(event);cb.on_state(event)
        self.assertEqual(len(web.status('cell')['articles']),1)

    def test_slow_client_has_bounded_queue_and_recovery_snapshot(self):
        q=web.bus.subscribe()
        for i in range(1005): web.bus.broadcast(dict(type='log',source='cell',data=str(i)))
        events=[]
        while not q.empty(): events.append(q.get_nowait())
        self.assertLessEqual(len(events),1000)
        self.assertEqual(events[0]['type'],'snapshot')
        web.bus.unsubscribe(q)

    def test_concurrent_start_only_one_worker(self):
        barrier=threading.Barrier(2)
        release=threading.Event()
        calls=[]
        def run(*args,**kwargs): calls.append(1); release.wait(5)
        def start():
            barrier.wait()
            try:
                web.start({'urls':['https://www.nature.com/nature/volumes/654/issues/8119']},'nature')
                return 200
            except web.HTTPException as exc: return exc.status_code
        with patch.object(web.SOURCES['nature'],'run_scraper',side_effect=run):
            try:
                with ThreadPoolExecutor(max_workers=2) as executor:
                    futures=[executor.submit(start) for _ in range(2)]
                    self.assertEqual(sorted(f.result() for f in futures),[200,409])
                thread=web.SESSIONS['nature'].thread
                self.assertEqual(len(calls),1)
            finally:
                release.set()
                thread=web.SESSIONS['nature'].thread
                if thread: thread.join(5)

    def test_cancel_takes_priority_over_cf_resume(self):
        session=web.SESSIONS['cell']
        session.stop_event.set()
        self.assertEqual(web.WebCallbacks('cell').cf_wait('url','url',1,5,lambda:True),'cancelled')

    def test_reset_export_guard_and_url_validation(self):
        web.SESSIONS['cell'].thread=object()
        for operation in (web.reset_cache,web.export_now):
            with self.assertRaises(web.HTTPException) as raised: operation('cell')
            self.assertEqual(raised.exception.status_code,409)
        with self.assertRaises(web.HTTPException): web.validate_urls(['http://localhost:8000'],'nature')
        with self.assertRaises(web.HTTPException) as raised: web.stop('invalid')
        self.assertEqual(raised.exception.status_code,400)


if __name__ == '__main__':
    unittest.main()
